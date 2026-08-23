"""Aggregation layer for the per-lot "Superficie" invoice table, modeled
after the uploaded reference invoice
(OKAPI_WCAG_LARAP_Invoice 1_24042026 v2.xlsx):

    N° | Lots | Superficie des terres inventoriées à 100% par OKAPI (ha)
       | Superficie des structures inventoriées par OKAPI (m²)

Unlike the existing per-PAP GNF monetary breakdown (facturation_page /
facturation_export in app.py, unchanged), this is a much simpler per-LOT
(num_batch) aggregation of RAW SURFACE AREA — no monetary/GNF figures at
all:
  - "Superficie des terres" = sum of ParcelleAgricole.superficieParcelle
    across all champs belonging to the lot, converted from m² to ha
    (/10000) to match the reference template's "ha" column.
  - "Superficie des structures" = sum of StructureItem.superficieSol
    across all structures belonging to the lot, in m² (no conversion,
    matches the reference template's "m²" column directly).

This module is purely additive: it does not touch or replace the existing
per-PAP GNF facturation logic in app.py / excel_export-style code.
"""


def build_lot_superficie_rows(champs, structures):
    """Returns (rows, totals) where `rows` is a list of dicts:
        {"num_batch": str, "superficie_terres_ha": float,
         "superficie_structures_m2": float}
    one row per distinct lot (num_batch) that has at least one champ or
    structure record, sorted by num_batch. `totals` is the same shape,
    summed across all rows (matching the reference template's Total row,
    which uses SUM formulas over C5:C8 and D5:D8)."""
    by_batch = {}

    for ch in champs:
        batch = ch.get("numBatch") or ""
        if not batch:
            continue
        entry = by_batch.setdefault(batch, {"superficie_terres_m2": 0.0, "superficie_structures_m2": 0.0})
        for p in ch.get("parcelles", []):
            entry["superficie_terres_m2"] += p.get("superficieParcelle", 0) or 0

    for s in structures:
        batch = s.get("numBatch") or ""
        if not batch:
            continue
        entry = by_batch.setdefault(batch, {"superficie_terres_m2": 0.0, "superficie_structures_m2": 0.0})
        for st in s.get("structures", []):
            entry["superficie_structures_m2"] += st.get("superficieSol", 0) or 0

    rows = []
    tot_terres_ha = 0.0
    tot_structures_m2 = 0.0
    for batch in sorted(by_batch.keys()):
        e = by_batch[batch]
        terres_ha = e["superficie_terres_m2"] / 10000.0
        structures_m2 = e["superficie_structures_m2"]
        rows.append({
            "num_batch": batch,
            "superficie_terres_ha": terres_ha,
            "superficie_structures_m2": structures_m2,
        })
        tot_terres_ha += terres_ha
        tot_structures_m2 += structures_m2

    totals = {
        "superficie_terres_ha": tot_terres_ha,
        "superficie_structures_m2": tot_structures_m2,
    }
    return rows, totals


def build_invoice_workbook(champs, structures):
    """Builds an openpyxl Workbook reproducing the exact layout of the
    uploaded reference invoice (OKAPI_WCAG_LARAP_Invoice 1_24042026 v2.xlsx):
    title rows, a 4-column header (N° / Lots / Superficie terres (ha) /
    Superficie structures (m²)), one data row per lot, and a Total row
    with SUM formulas — same structure, populated with live app data.
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side

    rows, totals = build_lot_superficie_rows(champs, structures)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice 1"

    thin = Side(style="thin")
    border_all = Border(top=thin, bottom=thin, left=thin, right=thin)
    font_normal = Font(name="Calibri", size=11)
    font_bold = Font(name="Calibri", size=11, bold=True)

    ws["A1"] = "OKAPI Environnement Conseil"
    ws["A1"].font = font_bold
    ws["A2"] = "Invoice 1"
    ws["A2"].font = font_bold

    header_row = 4
    headers = [
        "N°",
        "Lots",
        "Superficies des terres inventoriées à 100 % par OKAPI (ha)",
        "Superficies des structures inventoriées par OKAPI (m²)",
    ]
    for col_idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=label)
        cell.font = font_bold
        cell.border = border_all
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_start = header_row + 1
    row_idx = data_start
    for n, r in enumerate(rows, start=1):
        ws.cell(row=row_idx, column=1, value=n).border = border_all
        ws.cell(row=row_idx, column=1).font = font_normal
        ws.cell(row=row_idx, column=2, value=r["num_batch"]).border = border_all
        ws.cell(row=row_idx, column=2).font = font_normal
        c_terres = ws.cell(row=row_idx, column=3, value=round(r["superficie_terres_ha"], 6))
        c_terres.border = border_all
        c_terres.font = font_normal
        c_terres.number_format = "#,##0.00"
        c_struct = ws.cell(row=row_idx, column=4, value=(round(r["superficie_structures_m2"], 3) if r["superficie_structures_m2"] else None))
        c_struct.border = border_all
        c_struct.font = font_normal
        c_struct.number_format = "#,##0.00"
        row_idx += 1

    data_end = row_idx - 1

    # "Superficie totale à facturer" label row (matching reference row 9,
    # a plain label with no values of its own).
    ws.cell(row=row_idx, column=2, value="Superficie totale à facturer").font = font_normal
    row_idx += 1

    total_row = row_idx
    ws.cell(row=total_row, column=2, value="Total").font = font_bold
    ws.cell(row=total_row, column=2).alignment = Alignment(horizontal="right")
    ws.cell(row=total_row, column=2).border = border_all
    if data_end >= data_start:
        c_tot_terres = ws.cell(row=total_row, column=3, value=f"=SUM(C{data_start}:C{data_end})")
        c_tot_struct = ws.cell(row=total_row, column=4, value=f"=SUM(D{data_start}:D{data_end})")
    else:
        c_tot_terres = ws.cell(row=total_row, column=3, value=0)
        c_tot_struct = ws.cell(row=total_row, column=4, value=0)
    for c in (c_tot_terres, c_tot_struct):
        c.font = font_bold
        c.border = border_all
        c.number_format = "#,##0.00"

    ws.column_dimensions["A"].width = 5.45
    ws.column_dimensions["B"].width = 53.82
    ws.column_dimensions["C"].width = 23.54
    ws.column_dimensions["D"].width = 22.45

    return wb, rows, totals
