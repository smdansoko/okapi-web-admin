"""Generates the "Annuaire des sites de patrimoine culturel" Word (.docx)
export, modeled after the uploaded reference document (OKAPI_Annuaire des
sites Patrimoine culturel.docx): a title page, a summary table (VILLAGE /
ZONE / NOM DU SITE / TYPE / SOUS-TYPE / IMPORTANCE), then one heading +
detailed label/value table per site - built live from the `patrimoine_culturel`
survey records synced from the mobile app (see rapport_patrimoine_data.py).
"""
import io

from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def _set_cell_shading(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color})
    tc_pr.append(shd)


def _add_heading(doc, text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p


def _add_para(doc, text, size=10, justify=True):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.size = Pt(size)
    return p


def _add_table(doc, header, rows, col_widths_mm=None, font_size=8.5):
    n_cols = len(header)
    n_rows = 1 + len(rows)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for j, label in enumerate(header):
        cell = table.cell(0, j)
        cell.text = str(label)
        _set_cell_shading(cell, "EDEDED")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(font_size)

    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_size)

    if col_widths_mm:
        total = sum(col_widths_mm)
        usable_mm = 174.0
        for j, w in enumerate(col_widths_mm):
            width = Mm(usable_mm * (w / total))
            for row in table.rows:
                row.cells[j].width = width

    return table


def _title_page(doc, data):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("OKAPI ENVIRONNEMENT CONSEIL")
    run.bold = True
    run.font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ANNUAIRE DES SITES DE PATRIMOINE CULTUREL")
    run.bold = True
    run.font.size = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Inventaire du patrimoine culturel matériel et immatériel")
    run.font.size = Pt(12)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Nombre de sites recensés : {data['count']}")
    run.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Généré le {data['generated_at'].strftime('%d %B %Y à %H:%M')}")
    run.font.size = Pt(10)
    doc.add_page_break()


def _summary_table(doc, data):
    _add_heading(doc, "Tableau récapitulatif des sites de patrimoine culturel", size=13)
    header = ["VILLAGE", "ZONE", "NOM DU SITE", "TYPE", "SOUS-TYPE", "IMPORTANCE"]
    if not data["summary_rows"]:
        _add_para(doc, "Aucun site de patrimoine culturel n'a encore été synchronisé.", justify=False)
        return
    _add_table(doc, header, data["summary_rows"], col_widths_mm=[28, 30, 34, 30, 30, 22], font_size=9)
    doc.add_page_break()


def _site_section(doc, site, index, total):
    _add_heading(doc, site["heading"], size=12, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=8)

    n_rows = len(site["detail_rows"]) + 1
    table = doc.add_table(rows=n_rows, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    header_row = table.rows[0]
    header_row.cells[0].merge(header_row.cells[1])
    header_row.cells[0].text = site["nom_site"]
    _set_cell_shading(header_row.cells[0], "D9D2C5")
    for p in header_row.cells[0].paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(11)

    for i, (label, value) in enumerate(site["detail_rows"], start=1):
        c0 = table.cell(i, 0)
        c0.text = label
        _set_cell_shading(c0, "F2EFE9")
        for p in c0.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
        c1 = table.cell(i, 1)
        c1.text = value
        for p in c1.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)

    widths = [Mm(55), Mm(119)]
    for row in table.rows:
        row.cells[0].width = widths[0]
        row.cells[1].width = widths[1]

    if index < total - 1:
        doc.add_page_break()


def generate_rapport_patrimoine_docx(report_data=None):
    """Generates the "Annuaire des sites de patrimoine culturel" Word
    document as bytes."""
    if report_data is None:
        from rapport_patrimoine_data import build_patrimoine_report_data
        report_data = build_patrimoine_report_data()

    doc = Document()
    for section in doc.sections:
        section.left_margin = Mm(18)
        section.right_margin = Mm(18)

    _title_page(doc, report_data)
    _summary_table(doc, report_data)

    sites = report_data["sites"]
    for i, site in enumerate(sites):
        _site_section(doc, site, i, len(sites))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
