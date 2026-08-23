"""Generates the "Rapport" Word (.docx) export, modeled after the uploaded
reference document (OKAPI_AMCP_PARC_..._Rapport_provisoire_...docx): title
page, contrôle qualité roster, préambule narrative, résultats narrative,
and all 14 "Tableaux" as native Word tables (built live from the app's
synced data via rapport_data.py — the reference doc's Tableaux 5-14 were
embedded as unreadable EMF screenshots, so here they are native, editable
Word tables generated from the same live data instead).
"""
import io

from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from rapport_pdf import (
    EQUIPE_PARC,
    PREPARE_PAR,
)
from contract_pdf import fmt_number


def _set_cell_shading(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color})
    tc_pr.append(shd)


def _add_heading(doc, text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p


def _add_para(doc, text, size=10, justify=True):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.size = Pt(size)
    return p


def _add_table(doc, header, rows, total_row=None, col_widths_mm=None, font_size=8.5):
    n_cols = len(header)
    n_rows = 1 + len(rows) + (1 if total_row else 0)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for j, label in enumerate(header):
        cell = table.cell(0, j)
        cell.text = str(label)
        _set_cell_shading(cell, "EDEDED")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(font_size)

    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_size)

    if total_row:
        i = len(rows) + 1
        for j, val in enumerate(total_row):
            cell = table.cell(i, j)
            cell.text = str(val)
            _set_cell_shading(cell, "EDEDED")
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(font_size)

    if col_widths_mm:
        total = sum(col_widths_mm)
        usable_mm = 170.0  # approx usable width for A4 with normal margins
        for j, w in enumerate(col_widths_mm):
            width = Mm(usable_mm * (w / total))
            for row in table.rows:
                row.cells[j].width = width

    return table


def _roster_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    for i, (name, fonction) in enumerate(rows):
        c0 = table.cell(i, 0)
        c0.text = name
        for p in c0.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9.5)
        c1 = table.cell(i, 1)
        c1.text = fonction
        for p in c1.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9.5)
    return table


def _title_page(doc, data):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("L'ENVIRONNEMENT AU SERVICE DU DÉVELOPPEMENT DURABLE")
    run.bold = True
    run.font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PLAN D'ACTION DE RÉINSTALLATION ET DE COMPENSATION (PARC)")
    run.bold = True
    run.font.size = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("RAPPORT PROVISOIRE")
    run.bold = True
    run.font.size = Pt(14)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Lot(s) : {', '.join(data['batches']) if data['batches'] else '—'}")
    run.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Généré le {data['generated_at'].strftime('%d %B %Y à %H:%M')}")
    run.font.size = Pt(10)
    doc.add_page_break()


def _controle_qualite(doc):
    _add_heading(doc, "CONTRÔLE QUALITÉ", size=14)
    _add_heading(doc, "Préparé par", size=11, space_before=6, space_after=4)
    _roster_table(doc, PREPARE_PAR)
    _add_heading(doc, "Vérifié et approuvé par", size=11, space_before=10, space_after=4)
    _roster_table(doc, [("Guy RONDEAU", "Expert environnementaliste international, Directeur général")])
    _add_heading(doc, "Équipe PARC", size=11, space_before=10, space_after=4)
    _roster_table(doc, EQUIPE_PARC)
    doc.add_page_break()


def _preambule(doc, data):
    _add_heading(doc, "PRÉAMBULE", size=14)
    _add_para(
        doc,
        "Le Plan d'Action de Réinstallation et de Compensation (PARC) a été préparé par OKAPI Environnement "
        "Conseil SARL, un cabinet spécialisé dans l'accompagnement environnemental et social des projets "
        "d'infrastructure et d'exploitation minière en République de Guinée.",
    )
    _add_para(
        doc,
        "Ce document porte sur le dispositif de compensation lié aux emprises du Projet, permettant "
        "l'indemnisation des Personnes Affectées par le Projet (PAP) pour la perte de leurs terres et de leurs "
        "biens.",
    )
    _add_para(
        doc,
        "Le présent rapport présente l'ensemble des mesures techniques, sociales et financières mises en œuvre "
        "dans le cadre du processus de compensation, sur la base des données synchronisées à ce jour dans "
        f"l'application OKAPI Survey, soit {data['indemnisation_count']} enregistrement(s) d'indemnisation.",
    )
    _add_para(doc, "Ce rapport constitue également un état d'avancement de la mise en œuvre du PARC.")
    doc.add_page_break()


def _resultats_intro(doc):
    _add_heading(doc, "RÉSULTATS", size=14)
    _add_para(
        doc,
        "Ce rapport présente les activités d'acquisition foncière menées sur l'ensemble des lots synchronisés "
        "à ce jour dans l'application OKAPI Survey.",
    )


def _tableau1(doc, data):
    t1 = data["tableau1"]
    _add_heading(doc, "Tableau 1 : Accords signés / Localisation par type et nombre de responsables", size=10.5, space_before=8)
    header = ["Lot", "Ménage", "Lignage", "Collectif", "Masculin", "Féminin", "Total"]
    rows = [[r["lot"], r["menage"], r["lignage"], r["collectif"], r["masculin"], r["feminin"], r["total"]] for r in t1["rows"]]
    tot = t1["total"]
    total_row = ["TOTAL", tot["menage"], tot["lignage"], tot["collectif"], tot["masculin"], tot["feminin"], tot["total"]]
    _add_table(doc, header, rows, total_row, [40, 22, 22, 22, 22, 22, 20])


def _tableau2(doc, data):
    t2 = data["tableau2"]
    _add_heading(doc, "Tableau 2 : Récapitulatif des compensations", size=10.5, space_before=10)
    header = ["Lot", "Nombre de PAP", "Superficie (m²)", "Compensation (GNF)"]
    rows = [[r["lot"], r["nb_pap"], f"{r['superficie']:,.2f}".replace(",", " "), fmt_number(r["montant"])] for r in t2["rows"]]
    tot = t2["total"]
    total_row = ["TOTAL", tot["nb_pap"], f"{tot['superficie']:,.2f}".replace(",", " "), fmt_number(tot["montant"])]
    _add_table(doc, header, rows, total_row, [45, 30, 40, 45])


def _tableau3(doc, data):
    t3 = data["tableau3"]
    _add_heading(doc, "Tableau 3 : Récapitulatif des compensations par village", size=10.5, space_before=10)
    header = ["Lot", "Village", "Nombre de PAP", "Superficie (m²)", "Compensation (GNF)"]
    rows = [[r["lot"], r["village"], r["nb_pap"], f"{r['superficie']:,.2f}".replace(",", " "), fmt_number(r["montant"])] for r in t3["rows"]]
    tot = t3["total"]
    total_row = ["", "TOTAL", tot["nb_pap"], f"{tot['superficie']:,.2f}".replace(",", " "), fmt_number(tot["montant"])]
    _add_table(doc, header, rows, total_row, [30, 40, 28, 40, 42])


def _tableau4(doc, data):
    t4 = data["tableau4"]
    _add_heading(doc, "Tableau 4 : Récapitulatif des compensations des cultures annuelles", size=10.5, space_before=10)
    header = ["Culture", "Superficie (ha)", "Compensation (GNF)"]
    rows = [[r["culture"], f"{r['superficie']:,.4f}".replace(",", " "), fmt_number(r["montant"])] for r in t4["rows"]]
    tot = t4["total"]
    total_row = ["TOTAL", f"{tot['superficie']:,.4f}".replace(",", " "), fmt_number(tot["montant"])]
    _add_table(doc, header, rows, total_row, [60, 45, 45])
    doc.add_page_break()


def _tableau5(doc, data):
    t5 = data["tableau5"]
    _add_heading(doc, "Tableau 5 : Récapitulatif des compensations des arbres / lot", size=10.5, space_before=8)
    header = ["Lot", "Cultures pérennes (GNF)", "Espèces sauvages (GNF)", "Bois d'œuvre (GNF)", "Total (GNF)"]
    rows = [
        [r["lot"], fmt_number(r["cultures_perennes"]), fmt_number(r["especes_sauvages"]), fmt_number(r["bois_doeuvre"]), fmt_number(r["total"])]
        for r in t5["rows"]
    ]
    tb = t5["totals_by_cat"]
    total_row = ["TOTAL", fmt_number(tb["cultures_perennes"]), fmt_number(tb["especes_sauvages"]), fmt_number(tb["bois_doeuvre"]), fmt_number(t5["grand_total"])]
    _add_table(doc, header, rows, total_row, [28, 38, 38, 38, 32], font_size=7.5)


def _tableau6(doc, data):
    t6 = data["tableau6"]
    _add_heading(doc, "Tableau 6 : Récapitulatif des montants et superficie relatif au type de terrain / localisation par village", size=10.5, space_before=10)
    header = ["Type de terrain", "Village", "Superficie (m²)", "Montant (GNF)"]
    rows = [[r["type_terrain"], r["village"], f"{r['superficie']:,.2f}".replace(",", " "), fmt_number(r["montant"])] for r in t6["rows"]]
    tot = t6["total"]
    total_row = ["TOTAL", "", f"{tot['superficie']:,.2f}".replace(",", " "), fmt_number(tot["montant"])]
    _add_table(doc, header, rows, total_row, [55, 45, 40, 46], font_size=8.0)
    doc.add_page_break()


def _indemnisation_tables(doc, data, tableau_start=7):
    pages = data["indemnisation_pages"]
    header = [
        "N°", "Lot", "Prénom et Nom", "Code de l'enquête", "Genre",
        "Village", "Statut contrat", "Superficie (m²)", "Montant (GNF)",
    ]
    for idx, page_rows in enumerate(pages):
        n_tab = tableau_start + idx
        suffix = "" if idx == 0 else (" (suite et fin)" if idx == len(pages) - 1 else " (suite)")
        _add_heading(
            doc,
            f"Tableau {n_tab} : Tableau d'indemnisation des {data['indemnisation_count']} PAP{suffix}",
            size=10.5, space_before=8,
        )
        rows = [
            [
                r["n"], r["num_lot"], r["nom"], r["code_enquete"], r["genre"],
                r["village"], r["statut_contrat"],
                f"{r['superficie']:,.2f}".replace(",", " "), fmt_number(r["montant"]),
            ]
            for r in page_rows
        ]
        tot_sup = sum(r["superficie"] for r in page_rows)
        tot_montant = sum(r["montant"] for r in page_rows)
        total_row = ["", "", "", "", "", "", "TOTAL", f"{tot_sup:,.2f}".replace(",", " "), fmt_number(tot_montant)]
        _add_table(doc, header, rows, total_row, [10, 22, 45, 32, 16, 30, 26, 26, 30], font_size=7.5)
        if idx < len(pages) - 1:
            doc.add_page_break()


def generate_rapport_docx(menages, champs, structures, report_data=None):
    """Generates the Rapport Word document as bytes."""
    if report_data is None:
        from rapport_data import build_full_report_data
        report_data = build_full_report_data(menages, champs, structures)

    doc = Document()
    # Slightly narrower margins so wide tables fit better.
    for section in doc.sections:
        section.left_margin = Mm(18)
        section.right_margin = Mm(18)

    _title_page(doc, report_data)
    _controle_qualite(doc)
    _preambule(doc, report_data)
    _resultats_intro(doc)
    _tableau1(doc, report_data)
    _tableau2(doc, report_data)
    _tableau3(doc, report_data)
    _tableau4(doc, report_data)
    _tableau5(doc, report_data)
    _tableau6(doc, report_data)
    _indemnisation_tables(doc, report_data)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
