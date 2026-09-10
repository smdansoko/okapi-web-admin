"""Import the legacy KoboToolbox/ODK "SMB_Enquête ménages.xlsx" export into
the SMB project's SQLite database, producing records in the exact JSON shape
the Flutter mobile app / db.upsert_menage() already expect — mirroring
import_legacy_excel.py's approach for WCAG/SIMANDOU (Phase 4 pattern), but
for the SMB project ONLY.

Source: /home/user/uploaded_files/SMB_Enquête ménages.xlsx

The workbook has 2 sheets:
  - "SMB_ménages"     : one row per ménage (household)
  - "infos_individus" : one row per individu (household member), linked to
                        its parent ménage row via _parent_index == the
                        ménage row's _index (both are KoboToolbox/ODK
                        export bookkeeping columns). For this particular
                        export, each ménage has exactly ONE linked individu
                        (the chef de ménage / répondant), verified 1:1
                        (716 ménages <-> 716 individus, zero missing links).

Unlike the WCAG/SIMANDOU legacy exports, this SMB export has NO populated
`commune` column (always empty) and no separate GPS/coord columns filled
in. Instead:
  - Région / Préfecture / Sous-préfecture are FIXED for every SMB record,
    taken directly from SMB's own reference contract PDFs (Contrat
    ménage/lignage/communautaire (SMB).pdf, section "1. Identification"):
      Région = "Boké", Préfecture = "Boké", Sous-préfecture = "Dabiss"
    (all 3 reference PDFs consistently show this same triple for every
    SMB household sampled, confirming SMB's project area sits entirely
    within the Dabiss sous-préfecture of Boké préfecture/région).
  - District / Village: the Excel's `localite` column (e.g. "Sonsoli",
    "Boundhou-Lengué") is the full village/locality name (NOT a short
    code needing translation, unlike WCAG/SIMANDOU's `village` column) -
    used as-is for BOTH district and village, matching how WCAG/SIMANDOU
    ended up with district==village after their own code->name mapping.
  - The `village` column (e.g. "SS", "BO", "DJ") is a short internal
    tablet/Kobo code that duplicates `localite` 1:1 (verified: each code
    maps to exactly one `localite` value) - kept only as an internal
    sanity-check map (SMB_VILLAGE_MAP), not stored in the final record.

Run with:  python3 import_smb_menages.py [--dry-run]
"""

import argparse
import re
from collections import Counter
from datetime import datetime

import openpyxl

import db

SMB_XLSX = "/home/user/uploaded_files/SMB_Enquête ménages.xlsx"

# Fixed admin division for every SMB household, taken from SMB's own
# reference contract PDFs (100% consistent across all 3 contract types
# sampled: Ménage / Lignage / Communautaire).
SMB_REGION = "Boké"
SMB_PREFECTURE = "Boké"
SMB_SOUS_PREFECTURE = "Dabiss"


def _iso(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _read_sheet(ws):
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers) if h is not None}
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return idx, rows


def load_menage_workbook(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws_m = wb["SMB_ménages"]
    ws_i = wb["infos_individus"]

    midx, mrows = _read_sheet(ws_m)
    iidx, irows = _read_sheet(ws_i)

    mreal = [r for r in mrows if r[midx["code_menage"]]]
    ireal = [r for r in irows if r[iidx["id_individu"]]]

    index_counts = Counter(r[midx["_index"]] for r in mreal)
    duplicate_indices = {idx for idx, c in index_counts.items() if c > 1}
    if duplicate_indices:
        print(
            f"  ⚠️  Duplicate ménage _index values detected: {sorted(duplicate_indices)} "
            f"- falling back to id_individu-prefix matching for affected groups."
        )

    individus_by_parent_index = {}
    for r in ireal:
        pidx = r[iidx["_parent_index"]]
        individus_by_parent_index.setdefault(pidx, []).append(r)

    village_map = {}  # short code -> full locality name (sanity-check only)
    menages = []
    for mr in mreal:
        m_index = mr[midx["_index"]]
        code_menage = mr[midx["code_menage"]]
        ind_rows = individus_by_parent_index.get(m_index, [])
        if m_index in duplicate_indices:
            ind_rows = [
                r for r in ind_rows if str(r[iidx["id_individu"]]).startswith(code_menage + "-")
            ]

        village_code = mr[midx["village"]] or ""
        localite = mr[midx["localite"]] or ""
        if village_code:
            village_map[village_code] = localite

        individus = []
        for ir in ind_rows:
            num_ordre = ir[iidx["num_ordre_individu"]]
            individus.append(
                {
                    "id": ir[iidx["id_individu"]],
                    "numOrdreIndividu": num_ordre if num_ordre is not None else 1,
                    "nomPrenom": ir[iidx["nom_prenom"]] or "",
                    "sexe": ir[iidx["sexe"]] or "",
                    "relationCdm": ir[iidx["relation_cdm"]] or "",
                    "typeDePiece": ir[iidx["type_de_piece"]] or "",
                    "numeroPiece": ir[iidx["numero_piece"]] or "",
                    "dateEtablissementPiece": _iso(ir[iidx["date_etablissement_piece"]]),
                    "dateNaissance": _iso(ir[iidx["date_naissance"]]),
                    "telephone": ir[iidx["telephone"]] or "",
                    "situationMatrimoniale": ir[iidx["situation_matrimoniale"]] or "",
                    "nombreFemmes": ir[iidx["nombre_femmes"]],
                    "groupeEthnique": ir[iidx["groupe_ethnique"]] or "",
                    "autreGroupeEthnique": ir[iidx["autre_groupe_ethnique"]],
                    "nationalite": ir[iidx["nationalite"]] or "",
                    "autreNationalite": ir[iidx["autre_nationalite"]],
                    "handicap": ir[iidx["handicap"]] or "Non",
                    "handicapPrecis": ir[iidx["handicap_precis"]],
                    "photoProfilBase64": None,
                    "photoCniRectoBase64": None,
                    "photoCniVersoBase64": None,
                    "codePap": None,
                    # Legacy filename kept for the fallback photo directory
                    # linkage (see contract_pdf.py _fallback_photo_b64),
                    # matching the WCAG/SIMANDOU import convention exactly.
                    "photoMembreFilenameLegacy": ir[iidx["photo_membre"]],
                }
            )

        menages.append(
            {
                "mr": mr,
                "midx": midx,
                "individus": individus,
                "localite": localite,
                "n_individus_expected": len(ind_rows),
            }
        )

    return menages, village_map


def build_menage_record(entry):
    mr = entry["mr"]
    midx = entry["midx"]
    localite = entry["localite"]

    date_enquete = datetime.now().isoformat()  # no date_enquete column in this export
    repondant_cdm = mr[midx["repondant_cdm"]] or "Oui"

    record = {
        "id": mr[midx["code_menage"]],
        "dateEnquete": date_enquete,
        "enqueteurs": "",  # no enqueteurs column in this export
        "tablette": str(mr[midx["tablette"]] or ""),
        "numOrdreMenage": mr[midx["num_ordre_menage"]] or 1,
        "region": SMB_REGION,
        "prefecture": SMB_PREFECTURE,
        "sousPrefecture": SMB_SOUS_PREFECTURE,
        "district": localite,
        "village": localite,
        "codeMenage": mr[midx["code_menage"]],
        "residencePrincipale": mr[midx["residence_principale"]] or "Oui",
        "statutMenage": mr[midx["statut_menage"]] or "",
        "latitude": mr[midx.get("_coord_menage_latitude", -1)] if "_coord_menage_latitude" in midx else None,
        "longitude": mr[midx.get("_coord_menage_longitude", -1)] if "_coord_menage_longitude" in midx else None,
        "repondantCdm": repondant_cdm,
        "nomPrenomRepondant": mr[midx["nom_prenom_repondant"]] if "nom_prenom_repondant" in midx else None,
        "lienRepondantCdc": mr[midx["lien_repondant_cdc"]] if "lien_repondant_cdc" in midx else None,
        "telephoneRepondant": mr[midx["telephone_repondant"]] if "telephone_repondant" in midx else None,
        "typeDePieceRepondant": mr[midx["type_de_piece_repondant"]] if "type_de_piece_repondant" in midx else None,
        "numeroPieceRepondant": mr[midx["numero_piece_repondant"]] if "numero_piece_repondant" in midx else None,
        "datePieceRepondant": _iso(mr[midx["date_piece_repondant"]]) if "date_piece_repondant" in midx else None,
        "individus": entry["individus"],
        "createdAt": date_enquete,
        "updatedAt": date_enquete,
    }
    return record


def import_smb(dry_run=False):
    print(f"\n=== Importing SMB legacy data from {SMB_XLSX} ===")
    entries, village_map = load_menage_workbook(SMB_XLSX)
    print(f"  Ménages found (real rows): {len(entries)}")
    total_individus = sum(len(e["individus"]) for e in entries)
    print(f"  Individus found (real rows, grouped): {total_individus}")
    print(f"  Distinct village codes -> localité mapping (sanity check): {len(village_map)}")

    db.set_current_project("smb")
    existing_ids = {m["id"] for m in db.all_menages()}
    print(f"  Existing menages already in smb.db before import: {len(existing_ids)}")

    records = []
    collisions = []
    for entry in entries:
        rec = build_menage_record(entry)
        if rec["id"] in existing_ids:
            collisions.append(rec["id"])
        records.append(rec)

    print(f"  codeMenage collisions with existing data: {len(collisions)}")
    if collisions[:5]:
        print(f"    e.g. {collisions[:5]}")

    if dry_run:
        print("  --dry-run: no data written.")
        print(f"  Sample record: {records[0] if records else None}")
        return {
            "menages": len(records),
            "individus": total_individus,
            "collisions": len(collisions),
        }

    inserted = 0
    for rec in records:
        db.upsert_menage(rec, device_id="legacy-import-smb")
        inserted += 1

    counts = db.counts()
    print(f"  Done. {inserted} menages upserted. smb.db now has: {counts}")
    return {
        "menages": inserted,
        "individus": total_individus,
        "collisions": len(collisions),
        "counts_after": counts,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    db.init_db()
    result = import_smb(dry_run=args.dry_run)
    print("\n=== Summary ===")
    import json

    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
